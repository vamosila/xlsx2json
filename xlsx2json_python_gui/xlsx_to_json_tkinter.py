"""
Filename: xlsx_to_json_tkinter.py
Author: László Ádám Vámosi
Location: Córdoba, Spain
Date: 2025-11-15

Copyright (c) 2025 László Ádám Vámosi
All rights reserved.

---

Nombre de archivo: xlsx_to_json_tkinter.py
Autor: László Ádám Vámosi
Ubicación: Córdoba, España
Fecha: 2025-11-15

© 2025 László Ádám Vámosi
Todos los derechos reservados.
"""

import json
import math
import re
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
import os

def parse_datetime(date_str):
    match = re.search(r'(\d{2}/\d{2}/\d{4}) de (\d{2}:\d{2}) a (\d{2}:\d{2})', str(date_str))
    if not match:
        return None, None
    date, start_str, end_str = match.groups()
    start = datetime.strptime(f"{date} {start_str}", "%d/%m/%Y %H:%M")
    end = datetime.strptime(f"{date} {end_str}", "%d/%m/%Y %H:%M")
    if end < start:
        end += timedelta(days=1)
    return start, end

def round_duration(minutes, is_consecutive=False):
    if not is_consecutive and minutes < 60:
        minutes = 60
    elif minutes > 60:
        minutes = int(round(minutes / 10) * 10)
    return minutes

def round_total_duration(total_minutes):
    blocks = math.ceil(total_minutes / 30)
    return blocks * 30

def process_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    results = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        worker_id = row[0]
        worker_name = row[1]
        location = row[2]
        date_range = row[3]
        report_number = row[4]
        if not worker_id or not date_range:
            continue
        start, end = parse_datetime(date_range)
        if not start:
            continue
        duration_minutes = (end - start).total_seconds() / 60
        if worker_id not in results:
            results[worker_id] = {"worker_name": worker_name, "interventions": []}
        results[worker_id]["interventions"].append({
            "location": location,
            "start": start,
            "end": end,
            "report_number": report_number,
            "duration_minutes": duration_minutes
        })
    for worker_id, data in results.items():
        interventions = sorted(data["interventions"], key=lambda x: x["start"])
        total_minutes = 0
        processed = []
        for i, p in enumerate(interventions):
            is_consecutive = i > 0 and interventions[i - 1]["end"] == p["start"]
            original_minutes = p["duration_minutes"]
            adjusted_minutes = round_duration(original_minutes, is_consecutive)
            total_minutes += adjusted_minutes
            processed.append({
                "location": p["location"],
                "start": p["start"].strftime("%d/%m/%Y %H:%M"),
                "end": p["end"].strftime("%d/%m/%Y %H:%M"),
                "report_number": p["report_number"],
                "original_duration": f"{int(original_minutes // 60)}:{int(original_minutes % 60):02d}",
                "adjusted_duration": f"{int(adjusted_minutes // 60)}:{int(adjusted_minutes % 60):02d}"
            })
        rounded_total = round_total_duration(total_minutes)
        results[worker_id] = {
            "worker_name": data["worker_name"],
            "interventions": processed,
            "total_adjusted_hours": f"{int(rounded_total // 60)}:{int(rounded_total % 60):02d}"
        }
    return results

def select_file():
    current_dir = os.getcwd()
    file_path = filedialog.askopenfilename(initialdir=current_dir, filetypes=[("Excel files", "*.xlsx")])
    if not file_path:
        return
    data = process_excel(file_path)
    print(json.dumps(data, indent=4, ensure_ascii=False))
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_file = os.path.join(current_dir, f"{base_name}.json")
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)
    messagebox.showinfo("Hecho", f"Archivo JSON guardado como {output_file}")

root = tk.Tk()
root.title("Convertidor de Excel a JSON")
root.resizable(False, False)

window_width, window_height = 400, 150
screen_width = root.winfo_screenwidth()
screen_height = root.winfo_screenheight()
x = int((screen_width / 2) - (window_width / 2))
y = int((screen_height / 2) - (window_height / 2))
root.geometry(f"{window_width}x{window_height}+{x}+{y}")

tk.Label(root, text="Seleccione un archivo Excel para convertir a JSON").pack(pady=20)
tk.Button(root, text="Buscar archivo Excel", command=select_file, width=25).pack()

root.mainloop()
