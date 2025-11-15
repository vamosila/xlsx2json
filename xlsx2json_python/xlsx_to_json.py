"""
Filename: xlsx_to_json.py
Author: László Ádám Vámosi
Location: Córdoba, Spain
Date: 2025-11-15

Copyright (c) 2025 László Ádám Vámosi
All rights reserved.

---

Nombre de archivo: xlsx_to_json.py
Autor: László Ádám Vámosi
Ubicación: Córdoba, España
Fecha: 2025-11-15

© 2025 László Ádám Vámosi
Todos los derechos reservados.
"""

import json
import math
import re
from datetime import datetime, timedelta
from openpyxl import load_workbook
import os
import sys


# ------------------------------------------------------------
# DATE PARSER
# ------------------------------------------------------------
def parse_datetime(text):
    """
    Input example: '02/07/2025 de 23:15 a 00:30'
    Output: (datetime_start, datetime_end)
    """
    m = re.search(r"(\d{2}/\d{2}/\d{4}) de (\d{2}:\d{2}) a (\d{2}:\d{2})", str(text))
    if not m:
        return None, None

    date_str, start_str, end_str = m.groups()

    start = datetime.strptime(date_str + " " + start_str, "%d/%m/%Y %H:%M")
    end = datetime.strptime(date_str + " " + end_str, "%d/%m/%Y %H:%M")

    # crosses midnight
    if end < start:
        end += timedelta(days=1)

    return start, end


# ------------------------------------------------------------
# ROUNDING RULES
# ------------------------------------------------------------
def round_intervention(minutes, consecutive=False):
    """
    Rule 1: minimum 1 hour if NOT consecutive
    Rule 2: if >1h OR consecutive → round to nearest 10 minutes
    """
    if not consecutive and minutes < 60:
        return 60  # minimum of 1 hour

    return round(minutes / 10) * 10


def round_monthly(total_min):
    """
    Rule 3: monthly rounding up to 30-minute blocks (0.5 hours).
    """
    return math.ceil(total_min / 30) * 0.5


# ------------------------------------------------------------
# PROCESS THE EXCEL FILE
# ------------------------------------------------------------
def process_excel(file_path):

    # -------------------------------
    # ERROR HANDLING: missing file
    # -------------------------------
    if not os.path.exists(file_path):
        print(f"ERROR: The Excel file '{file_path}' does not exist.")
        sys.exit(1)

    try:
        wb = load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"ERROR: Cannot open Excel file '{file_path}'.")
        print("Reason:", e)
        sys.exit(1)

    ws = wb.active
    workers = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        worker_id = row[0]
        name = row[1]
        location = row[2]
        date_range = row[3]
        report_number = row[4]

        if not worker_id or not date_range:
            continue

        start, end = parse_datetime(date_range)
        if not start:
            continue

        duration = (end - start).total_seconds() / 60

        if worker_id not in workers:
            workers[worker_id] = {
                "worker_name": name,
                "rows": []
            }

        workers[worker_id]["rows"].append({
            "start": start,
            "end": end,
            "location": location,
            "report_number": report_number,
            "duration_min": duration
        })

    results = {}

    for worker_id, data in workers.items():
        rows = sorted(data["rows"], key=lambda x: x["start"])
        interventions = []
        total_adjusted = 0

        for i, r in enumerate(rows):
            original_minutes = r["duration_min"]

            consecutive = (i > 0 and rows[i - 1]["end"] == r["start"])

            adjusted_minutes = round_intervention(original_minutes, consecutive)
            total_adjusted += adjusted_minutes

            interventions.append({
                "location": r["location"],
                "start": r["start"].strftime("%d/%m/%Y %H:%M"),
                "end": r["end"].strftime("%d/%m/%Y %H:%M"),
                "report_number": r["report_number"],
                "original_duration": f"{int(original_minutes//60)}:{int(original_minutes%60):02d}",
                "adjusted_duration": f"{int(adjusted_minutes//60)}:{int(adjusted_minutes%60):02d}",
                "consecutive_with_previous": consecutive
            })

        final_hours = round_monthly(total_adjusted)

        results[worker_id] = {
            "worker_name": data["worker_name"],
            "interventions": interventions,
            "total_adjusted_hours": final_hours
        }

    return results


# ------------------------------------------------------------
# MAIN PROGRAM
# ------------------------------------------------------------
if __name__ == "__main__":
    input_file = "Example.xlsx"
    output_file = "Example.json"

    data = process_excel(input_file)

    # Save JSON
    try:
        with open(output_file, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
    except Exception as e:
        print("ERROR: Cannot save JSON file.")
        print("Reason:", e)
        sys.exit(1)

    print("\n============== CONSOLE OUTPUT ==============\n")

    for worker_id, info in data.items():
        print(f"--- Worker {worker_id} / {info['worker_name']} ---")
        print(f"Total adjusted hours (monthly rounding): {info['total_adjusted_hours']} h\n")

        for i, interv in enumerate(info["interventions"], 1):
            print(f"  Intervention {i}:")
            print(f"    Location: {interv['location']}")
            print(f"    Start: {interv['start']}")
            print(f"    End: {interv['end']}")
            print(f"    Report No.: {interv['report_number']}")
            print(f"    Original duration: {interv['original_duration']}")
            print(f"    Adjusted duration: {interv['adjusted_duration']}")
            print(f"    Consecutive with previous: {interv['consecutive_with_previous']}")
            print()

        print("-------------------------------------------\n")

    print(f"JSON saved as: {output_file}")
